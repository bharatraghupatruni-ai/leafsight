"""
LEAFSIGHT — Complete Final Academic Project Report Generator
Generates:
  docs/LEAFSIGHT_Final_Project_Report.docx
  docs/LEAFSIGHT_Final_Project_Report.pdf
"""

import os
import sys
from pathlib import Path
import docx
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import parse_xml, OxmlElement
from docx.oxml.ns import nsdecls, qn
import win32com.client


# ==============================================================================
# STYLE UTILITIES
# ==============================================================================

COLOR_NAVY = RGBColor(27, 54, 93)      # #1B365D - Primary headings
COLOR_GREEN = RGBColor(22, 97, 62)     # #16613E - Subheadings & Accents
COLOR_DARK = RGBColor(34, 34, 34)      # #222222 - Body text
COLOR_MUTED = RGBColor(100, 110, 115)  # #646E73 - Captions & Footnotes

HEX_NAVY = "1B365D"
HEX_GREEN = "16613E"
HEX_LIGHT_BG = "F4F6F9"
HEX_ALT_ROW = "F9FAFB"
HEX_BORDER = "D0D7DE"


def set_cell_background(cell, fill_hex):
    tcPr = cell._element.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tcPr.append(shd)


def set_cell_margins(cell, top=100, bottom=100, left=140, right=140):
    tcPr = cell._element.get_or_add_tcPr()
    tcMar = parse_xml(
        f'<w:tcMar {nsdecls("w")}>'
        f'<w:top w:w="{top}" w:type="dxa"/>'
        f'<w:bottom w:w="{bottom}" w:type="dxa"/>'
        f'<w:left w:w="{left}" w:type="dxa"/>'
        f'<w:right w:w="{right}" w:type="dxa"/>'
        f'</w:tcMar>'
    )
    tcPr.append(tcMar)


def set_table_borders(table, color=HEX_BORDER):
    tblPr = table._element.xpath('w:tblPr')
    if tblPr:
        borders = parse_xml(
            f'<w:tblBorders {nsdecls("w")}>'
            f'<w:top w:val="single" w:sz="6" w:space="0" w:color="{color}"/>'
            f'<w:bottom w:val="single" w:sz="6" w:space="0" w:color="{color}"/>'
            f'<w:insideH w:val="single" w:sz="4" w:space="0" w:color="{color}"/>'
            f'<w:insideV w:val="none"/>'
            f'<w:left w:val="none"/>'
            f'<w:right w:val="none"/>'
            f'</w:tblBorders>'
        )
        tblPr[0].append(borders)


def format_table(table, col_widths, headers, data, align_cols=None):
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(table)

    # Header Row
    hdr_cells = table.rows[0].cells
    for i, title in enumerate(headers):
        hdr_cells[i].text = title
        set_cell_background(hdr_cells[i], HEX_NAVY)
        set_cell_margins(hdr_cells[i], top=120, bottom=120, left=140, right=140)
        p = hdr_cells[i].paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        for r in p.runs:
            r.font.name = 'Arial'
            r.font.size = Pt(9.5)
            r.font.bold = True
            r.font.color.rgb = RGBColor(255, 255, 255)

    # Data Rows
    for r_idx, row_data in enumerate(data):
        row = table.rows[r_idx + 1]
        bg_color = HEX_ALT_ROW if (r_idx % 2 == 1) else "FFFFFF"
        for c_idx, val in enumerate(row_data):
            cell = row.cells[c_idx]
            cell.text = str(val)
            set_cell_background(cell, bg_color)
            set_cell_margins(cell, top=80, bottom=80, left=120, right=120)
            p = cell.paragraphs[0]
            if align_cols and c_idx in align_cols:
                p.alignment = align_cols[c_idx]
            else:
                p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            for r in p.runs:
                r.font.name = 'Arial'
                r.font.size = Pt(9)
                r.font.color.rgb = COLOR_DARK

    # Column widths
    for row in table.rows:
        for idx, width in enumerate(col_widths):
            row.cells[idx].width = Inches(width)


def add_callout(doc, text, title="NOTE"):
    tbl = doc.add_table(rows=1, cols=1)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    c = tbl.cell(0, 0)
    c.width = Inches(6.5)
    set_cell_background(c, HEX_LIGHT_BG)
    set_cell_margins(c, top=100, bottom=100, left=160, right=160)
    tcPr = c._element.get_or_add_tcPr()
    borders = parse_xml(
        f'<w:tcBorders {nsdecls("w")}>'
        f'<w:left w:val="single" w:sz="24" w:space="0" w:color="{HEX_NAVY}"/>'
        f'<w:top w:val="none"/><w:bottom w:val="none"/><w:right w:val="none"/>'
        f'</w:tcBorders>'
    )
    tcPr.append(borders)
    p = c.paragraphs[0]
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.line_spacing = 1.15
    r_lbl = p.add_run(f"{title}: ")
    r_lbl.bold = True
    r_lbl.font.name = 'Arial'
    r_lbl.font.size = Pt(9.5)
    r_lbl.font.color.rgb = COLOR_NAVY
    r_txt = p.add_run(text)
    r_txt.font.name = 'Arial'
    r_txt.font.size = Pt(9.5)
    r_txt.font.color.rgb = COLOR_DARK
    doc.add_paragraph().paragraph_format.space_after = Pt(4)


def add_code_snippet(doc, code_str, caption=None):
    if caption:
        p_cap = doc.add_paragraph()
        p_cap.paragraph_format.space_before = Pt(6)
        p_cap.paragraph_format.space_after = Pt(2)
        r = p_cap.add_run(caption)
        r.bold = True
        r.font.name = 'Arial'
        r.font.size = Pt(9)
        r.font.color.rgb = COLOR_NAVY

    tbl = doc.add_table(rows=1, cols=1)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    c = tbl.cell(0, 0)
    c.width = Inches(6.5)
    set_cell_background(c, "F8F9FA")
    set_cell_margins(c, top=90, bottom=90, left=130, right=130)
    tcPr = c._element.get_or_add_tcPr()
    borders = parse_xml(
        f'<w:tcBorders {nsdecls("w")}>'
        f'<w:left w:val="single" w:sz="6" w:space="0" w:color="{HEX_BORDER}"/>'
        f'<w:top w:val="single" w:sz="6" w:space="0" w:color="{HEX_BORDER}"/>'
        f'<w:bottom w:val="single" w:sz="6" w:space="0" w:color="{HEX_BORDER}"/>'
        f'<w:right w:val="single" w:sz="6" w:space="0" w:color="{HEX_BORDER}"/>'
        f'</w:tcBorders>'
    )
    tcPr.append(borders)
    p = c.paragraphs[0]
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.line_spacing = 1.05
    r = p.add_run(code_str.strip())
    r.font.name = "Consolas"
    r.font.size = Pt(8.5)
    r.font.color.rgb = RGBColor(35, 35, 35)
    doc.add_paragraph().paragraph_format.space_after = Pt(4)


def add_figure(doc, img_path, caption, width=5.5):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(8)
    p.paragraph_format.space_after = Pt(2)
    if os.path.exists(img_path):
        p.add_run().add_picture(img_path, width=Inches(width))
    else:
        r = p.add_run(f"[Figure missing: {img_path}]")
        r.font.color.rgb = RGBColor(200, 0, 0)

    p_cap = doc.add_paragraph()
    p_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_cap.paragraph_format.space_before = Pt(2)
    p_cap.paragraph_format.space_after = Pt(10)
    r_cap = p_cap.add_run(caption)
    r_cap.font.name = 'Arial'
    r_cap.font.size = Pt(9)
    r_cap.font.italic = True
    r_cap.font.color.rgb = COLOR_MUTED


def add_p(doc, text, bold_prefix=None, space_after=6, italic=False):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.line_spacing = 1.15
    if bold_prefix:
        r_pre = p.add_run(bold_prefix)
        r_pre.bold = True
        r_pre.font.name = 'Arial'
        r_pre.font.size = Pt(10.5)
        r_pre.font.color.rgb = COLOR_DARK
    r_txt = p.add_run(text)
    r_txt.font.name = 'Arial'
    r_txt.font.size = Pt(10.5)
    r_txt.font.italic = italic
    r_txt.font.color.rgb = COLOR_DARK
    return p


def add_bullet(doc, text, bold_prefix=None):
    p = doc.add_paragraph(style='List Bullet')
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.line_spacing = 1.15
    if bold_prefix:
        r_pre = p.add_run(bold_prefix)
        r_pre.bold = True
        r_pre.font.name = 'Arial'
        r_pre.font.size = Pt(10)
        r_pre.font.color.rgb = COLOR_DARK
    r_txt = p.add_run(text)
    r_txt.font.name = 'Arial'
    r_txt.font.size = Pt(10)
    r_txt.font.color.rgb = COLOR_DARK
    return p


def add_h1(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(18)
    p.paragraph_format.space_after = Pt(8)
    p.paragraph_format.keep_with_next = True
    r = p.add_run(text)
    r.font.name = 'Arial'
    r.font.size = Pt(16)
    r.font.bold = True
    r.font.color.rgb = COLOR_NAVY
    return p


def add_h2(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(14)
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.keep_with_next = True
    r = p.add_run(text)
    r.font.name = 'Arial'
    r.font.size = Pt(13)
    r.font.bold = True
    r.font.color.rgb = COLOR_GREEN
    return p


def add_h3(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(10)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.keep_with_next = True
    r = p.add_run(text)
    r.font.name = 'Arial'
    r.font.size = Pt(11)
    r.font.bold = True
    r.font.color.rgb = COLOR_NAVY
    return p


print("Helper utilities loaded successfully.")
